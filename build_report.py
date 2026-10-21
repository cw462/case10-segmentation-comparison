"""Summarize completed native-grid Case10 evaluations without inventing GT scores."""
import csv
import hashlib
import json
import math
import unicodedata
from datetime import datetime
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(sys.argv[1]).resolve()
RES = ROOT / 'results'

def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))

def sha256(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()

prep = read_json(RES / 'preparation.json')
manifest = read_json(ROOT / 'tool_manifest.json')
lineage = read_json(RES / 'lineage_summary.json')
versions = read_json(RES / 'environment_versions.json')
names = {'pulmonary_artery':'肺动脉', 'pulmonary_vein':'肺静脉', 'airway':'气道', 'lung':'肺'}
reports = {s:read_json(RES / (s+'_metrics.json')) for s in names}
assert prep['complete'] and lineage['complete'] and all(r['complete'] for r in reports.values())

# Input/package integrity and checks independent of the exported spreadsheet.
checks = []
def check(label, okay):
    checks.append({'check':label, 'passed':bool(okay)})
    if not okay:
        raise AssertionError(label)

check('source ZIP SHA256 unchanged', sha256(manifest['archive']) == manifest['archive_sha256'])
for item in manifest['files']:
    check('unmodified tool: '+item['file'], sha256(ROOT/'tool'/item['file']) == item['sha256'])
for s, p in prep['structures'].items():
    for pathkey, hashkey in [('original_before_file','before_source_sha256'),
                             ('original_after_file','after_source_sha256'),
                             ('before_file','before_copy_sha256'),('after_file','after_nifti_sha256')]:
        check(s+' '+pathkey+' SHA256 unchanged', sha256(p[pathkey]) == p[hashkey])
    check(s+' count partition', p['shared_voxels']+p['added_voxels']==p['after_voxels'] and p['shared_voxels']+p['removed_voxels']==p['before_voxels'])
    air=reports[s]['metric_sets']['airrc']['metrics']
    check(s+' Dice from counts', math.isclose(air[s+'.DSC'], 2*p['shared_voxels']/(p['before_voxels']+p['after_voxels']), abs_tol=1e-12))
    check(s+' precision from counts', math.isclose(air[s+'.Precision'],p['shared_voxels']/p['before_voxels'],abs_tol=1e-12))
    check(s+' recall from counts', math.isclose(air[s+'.Recall'],p['shared_voxels']/p['after_voxels'],abs_tol=1e-12))
    for connectivity in ['6','26']:
        for stage in ['before','after']:
            cc=p[stage+'_components'][connectivity]
            check(s+' '+stage+' '+connectivity+' CC count',cc['largest_voxels']+cc['outside_largest_voxels']==p[stage+'_voxels'])
check('independent GT not claimed', all(r['independent_gt_available'] is False for r in reports.values()))

notes = [
 ('主要结论','此次修改集中于肺动脉和肺静脉；气道、肺掩码逐体素完全相同。血管主块占比提高、非主块体素减少，同时引入跨类别重叠。'),
 ('比较方向','版本一致性表以补标后的掩码为参考、补标前为待比较对象。反向计算时 Precision / Recall 会互换。'),
 ('不能推断的结论','没有独立于两版本的真值，不能给出补标前后对同一独立真值的准确率提升。未把补标后与自身比较的100%当作性能提升。'),
 ('工具来源','直接调用用户 ZIP 中的 VesMamba / HiPaS / AirRC 指标函数；原包26个文件保持原样，26项自带测试通过。包本身属于独立实现，不等同于已经验证的论文官方评测器。'),
 ('坐标与范围','512×512×657；体素间距0.8769530058×0.8769530058×0.625 mm；NRRD LPS 转为 NIfTI RAS 后与原标签网格一致，无重采样、平滑、删碎片或优先级合并。混淆计数使用整幅原始体积。'),
 ('互相重叠的标签','各结构保留为独立二值通道。补标后动脉/静脉836体素、气道/动脉27体素重叠，未通过覆盖标签人为消除。'),
 ('连通性定义','主表使用26邻域，面/边/角接触均算连通；另列6邻域，只有面接触算连通。组件数降低和主块占比提高仅说明几何连通性变化。'),
 ('连通块去向','按原组件的保留体素追踪。保留体素全部进入新主块不代表原组件完整保留；组件减少也可能来自删除。'),
 ('Dice与clDice','Dice描述体积重叠；clDice描述骨架与掩码的相互覆盖。此处均是前后版本一致性，不能用其差值直接判定补标对错。'),
 ('HD95定义','surface-distance 0.1：两方向按表面元面积加权的95%距离分位数，取较大值，单位mm。局部修改占比较小时可为0，不表示所有表面完全相同。'),
 ('NSD阈值','本次固定1.0 mm用于探索性比较，不宣称该阈值为论文原始协议。表面元面积加权。'),
 ('BC_proxy','Lee骨架中26邻域度数≥3的骨架体素形成的26连通簇数量，称分叉簇代理计数；不是人工确认的真实分支数，也不是分支检出率。'),
 ('SL_voxels_proxy','Lee骨架前景体素数；单位为体素，不是毫米长度。两版在同一原生网格比较；未采用原论文的标准化网格或骨架方法。'),
 ('代理指标的增减','BC / SL 增加描述骨架数量变化，不能单独证明分割更准确或分支召回提高。'),
 ('HiPaS调用适配','对真实二值动脉、静脉通道分别调用，保留对应类别结果；弃去为接口兼容设置的空类别结果及MCS。原包协议字段与适配说明一并保存在JSON。'),
 ('AUC缺失','未提供概率图，NA；不使用二值标签代替概率。'),
 ('BDR与DLR缺失','缺少经过确认并冻结的参考分支ID中心线图，NA；不从补标后自动造出真值再声称独立分支检出率。'),
 ('MCS缺失','补标后动静脉不互斥，标准单标签MCS不适用，NA；需先人工确定重叠归属。'),
 ('肺与气道壁','肺只提供辅助二值重叠与结构变化；肺不是气道壁。肺未运行血管/气道骨架与表面指标。'),
 ('统计范围','单个病例的描述性比较；不提供群体显著性结论、置信区间或未经支持的准确率增益。'),
 ('缺失值约定','NA表示本次未计算/条件不足，原因见本说明；空值不按0解释。'),
 ('源数据','原始与补标源文件未改动；计算副本、转换后NIfTI、增删掩码、原始JSON、工具哈希与复跑脚本均保存在本目录。'),
 ('前版本来源',prep['before_source']), ('后版本来源',prep['after_source']),
 ('ZIP来源',manifest['archive']), ('ZIP SHA256',manifest['archive_sha256']),
 ('运行环境',versions['python']+'; '+json.dumps(versions['packages'],ensure_ascii=False)),
]

wb = Workbook()
wb.remove(wb.active)
tables={}
def sheet(title, headers, rows, widths=None, formats=None):
    ws=wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    ws.sheet_view.showGridLines=False
    for c in ws[1]:
        c.fill=PatternFill('solid',fgColor='17365D')
        c.font=Font(name='微软雅黑',bold=True,color='FFFFFF',size=11)
        c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=36
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='微软雅黑',size=10)
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if c.row%2==0:
                c.fill=PatternFill('solid',fgColor='EEF3F8')
            if isinstance(c.value,(float,int)):
                c.number_format='#,##0.0000' if isinstance(c.value,float) else '#,##0'
            if formats and c.column in formats and isinstance(c.value,(float,int)):
                c.number_format=formats[c.column]
    for idx in range(1,len(headers)+1):
        ws.column_dimensions[get_column_letter(idx)].width=(widths or {}).get(idx,22)
    for row in ws.iter_rows(min_row=2):
        lines=1
        for c in row:
            if isinstance(c.value,str):
                display_length=sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in c.value)
                width=ws.column_dimensions[get_column_letter(c.column)].width
                lines=max(lines,math.ceil(display_length/max(width-2,1)))
        ws.row_dimensions[row[0].row].height=max(30,lines*15)
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1'
    tables[title]={'headers':headers,'rows':rows}
    return ws

overview=sheet('说明与结论',['项目','内容'],notes,{1:25,2:118})
for i in range(2,overview.max_row+1):
    overview.row_dimensions[i].height=44 if len(str(overview.cell(i,2).value))>90 else 32

structrows=[]
for s,name in names.items():
    p=prep['structures'][s]
    for label,b,a,unit,note in [
        ('前景体素数',p['before_voxels'],p['after_voxels'],'体素',''),
        ('体积',p['before_volume_ml'],p['after_volume_ml'],'mL',''),
        ('26邻域连通块总数',p['before_components']['26']['components'],p['after_components']['26']['components'],'块','包含最大主块'),
        ('26邻域主块占比',100*p['before_components']['26']['largest_fraction'],100*p['after_components']['26']['largest_fraction'],'%','差值单位为百分点；不等同于解剖正确率'),
        ('26邻域主块以外体素',p['before_components']['26']['outside_largest_voxels'],p['after_components']['26']['outside_largest_voxels'],'体素',''),
        ('26邻域≥20体素的非主块',p['before_components']['26']['components_ge20']-1,p['after_components']['26']['components_ge20']-1,'块','排除最大主块'),
        ('26邻域≥100体素的非主块',p['before_components']['26']['components_ge100']-1,p['after_components']['26']['components_ge100']-1,'块','排除最大主块'),
        ('6邻域连通块总数',p['before_components']['6']['components'],p['after_components']['6']['components'],'块','仅面接触算连通'),
        ('6邻域主块占比',100*p['before_components']['6']['largest_fraction'],100*p['after_components']['6']['largest_fraction'],'%','差值单位为百分点'),
    ]:
        relative=(a-b)/b if b and unit!='%' else 'NA'
        if s=='lung' and '主块' in label:
            note+='；肺26邻域可经对角接触连为大块，不能据此判断左右肺解剖连通'
        structrows.append([name,label,b,a,a-b,unit,relative,note])
sheet('真实结构变化',['结构','指标','补标前','补标后','后减前','数值单位','相对变化','说明'],structrows,{2:32,8:72},{7:'0.00%'})

editrows=[]
agreement=[]
for s,name in names.items():
    p=prep['structures'][s]
    editrows.append([name,p['added_voxels'],p['removed_voxels'],p['net_voxel_change'],p['net_volume_change_ml'],p['shared_voxels'],'是' if p['identical_voxelwise'] else '否'])
    r=reports[s];v=r['metric_sets']['vesmamba'].get('metrics',{});air=r['metric_sets']['airrc']['metrics'];h=r['metric_sets'].get('hipas',{}).get('metrics',{})
    prefix='artery' if s=='pulmonary_artery' else 'vein'
    agreement.append([name,air[s+'.DSC'],v.get('binary.clDice','NA'),v.get('binary.HD95_mm','NA'),v.get('binary.NSD','NA'),air[s+'.Precision'],air[s+'.Recall'],air[s+'.FPR'],h.get(prefix+'_SPE','NA'),'补标后版本；非独立GT'])
sheet('增加删除体素',['结构','新增体素','删除体素','净增体素','净体积变化_mL','两版共有体素','逐体素完全相同'],editrows)
sheet('前后版本一致性',['结构','Dice_DSC','clDice','HD95_mm','NSD_1mm','Precision','Recall_SEN','FPR','SPE','参考对象'],agreement,{10:38},{2:'0.00%',3:'0.00%',4:'0.000',5:'0.00%',6:'0.00%',7:'0.00%',8:'0.000000%',9:'0.000000%'})

proxyrows=[]
for s in ['pulmonary_artery','pulmonary_vein']:
    for metric,d in reports[s]['actual_proxy_changes'].items():
        proxyrows.append([names[s],metric,'分叉簇代理数' if metric=='BC_proxy' else '骨架体素数',d['before'],d['after'],d['delta_after_minus_before'],d['relative_change_percent']/100,'同网格Lee骨架代理；增多不等于准确率或分支召回提高'])
sheet('骨架代理变化',['结构','代码指标','含义','补标前','补标后','后减前','相对变化','解释限制'],proxyrows,{2:28,8:65},{4:'#,##0',5:'#,##0',6:'+#,##0;-#,##0;0',7:'0.00%'})

overlaprows=[]
for pair,b in prep['overlaps']['before'].items():
    a=prep['overlaps']['after'][pair]
    overlaprows.append([' / '.join(names[x] for x in pair.split('__')),b['voxels'],a['voxels'],a['voxels']-b['voxels'],b['components_26'],a['components_26'],a['volume_ml']])
sheet('类别重叠',['类别对','补标前重叠体素','补标后重叠体素','新增重叠体素','前重叠区域数_26','后重叠区域数_26','后重叠体积_mL'],overlaprows,{1:28})

linerows=[]
for s,ls in lineage['structures'].items():
    for key,label in [('retained_only_in_new_main','原非主块有保留体素且保留部分全归入新主块'),('all_deleted','原非主块全部体素被删除'),('retained_only_in_one_new_nonmain','原非主块保留部分仍处于单个非主块'),('retained_in_multiple_new_components','原非主块保留部分散布于多个新块')]:
        linerows.append([names[s],label,ls['original_nonmain_components']['categories'][key],'块','以上4项划分原始非主块；归入主块不代表原组件完整保留'])
    for key,label in [('total','原非主块体素总数'),('retained_in_after_main','原非主块体素现归入主块'),('retained_in_after_nonmain','原非主块体素现仍在非主块'),('deleted','原非主块体素被删除')]:
        linerows.append([names[s],label,ls['original_nonmain_voxels'][key],'体素','后三项之和等于原非主块体素总数'])
    added=ls['after_components_containing_only_added_voxels']
    linerows.extend([[names[s],'纯新增体素形成的新孤立块',len(added),'块','大小(体素): '+', '.join(str(x['voxels']) for x in sorted(added,key=lambda x:-x['voxels']))],
                     [names[s],'纯新增孤立块体素数',sum(x['voxels'] for x in added),'体素','几何追踪；不判定解剖正确性']])
sheet('连通块去向',['结构','原组件或体素的去向','数量','单位','说明'],linerows,{2:58,5:80})

rawrows=[]
for s,r in reports.items():
    for group,g in r['metric_sets'].items():
        for key,val in g.get('metrics',{}).items():
            rawrows.append([names[s],group,key,val if val is not None else 'NA',g.get('status',{}).get(key,''),'pred=补标前；gt=补标后；非独立真值'])
sheet('工具原始指标',['结构','工具指标组','代码指标名','原始值_比例为0至1','状态','说明'],rawrows,{2:22,3:48,4:28,5:52,6:57})

regionfile=RES/'overlap_regions.csv'
if regionfile.exists():
    with regionfile.open(encoding='utf-8-sig',newline='') as f:
        reader=csv.reader(f);headers=next(reader);rows=list(reader)
    for row in rows:
        for index in (1,2,3):
            row[index]=int(row[index])
        for index in (4,5,6,7):
            row[index]=float(row[index])
    sheet('重叠区域定位',headers,rows,{1:38,9:48,10:50,11:100})

for title,data in tables.items():
    with (RES/(title+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(data['headers']);writer.writerows(data['rows'])

summary={
    'created':datetime.now().isoformat(timespec='seconds'),
    'independent_ground_truth_available':False,
    'reference_role':'supplemented version; version agreement only',
    'preparation':prep,'metrics':reports,'component_lineage':lineage,
    'notes_zh':dict(notes),'tool_manifest':manifest,'versions':versions,
    'verification':checks,
}
(RES/'comparison_report.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')

bookfile=ROOT/'Case10补标前后指标对比.xlsx'
wb.save(bookfile)
loaded=load_workbook(bookfile,data_only=True,read_only=True)
check('workbook sheet count', len(loaded.sheetnames)==len(tables))
for title,table in tables.items():
    check('workbook rows: '+title, loaded[title].max_row==len(table['rows'])+1)
for rowidx,s in enumerate(names,start=2):
    check('workbook Dice: '+s, math.isclose(loaded['前后版本一致性'].cell(rowidx,2).value,prep['structures'][s]['version_agreement_Dice'],abs_tol=1e-12))
loaded.close()

lines=[
 'Case10 补充标注前后指标比较',
 '生成时间：'+summary['created'],
 '',
 '已使用指定ZIP的原始指标代码完成比较。原数据未修改。',
 '最主要的结论：气道与肺逐体素完全未变；修改集中于动脉和静脉，主块占比提高、非主块体素减少，同时出现新的类别重叠。',
 '',
 '一、真实前后变化（26邻域）',
]
for s in ['pulmonary_artery','pulmonary_vein']:
    p=prep['structures'][s];b=p['before_components']['26'];a=p['after_components']['26']
    lines.extend([
        names[s]+':',
        f"  体积 {p['before_volume_ml']:.3f} -> {p['after_volume_ml']:.3f} mL，净增 {p['net_volume_change_ml']:.3f} mL。",
        f"  新增 {p['added_voxels']:,} 体素，删除 {p['removed_voxels']:,} 体素。",
        f"  连通块 {b['components']} -> {a['components']}；主块占比 {100*b['largest_fraction']:.4f}% -> {100*a['largest_fraction']:.4f}%。",
        f"  主块外体素 {b['outside_largest_voxels']:,} -> {a['outside_largest_voxels']:,}。",
    ])
    for measure,d in reports[s]['actual_proxy_changes'].items():
        lines.append(f"  {measure}: {d['before']:,.0f} -> {d['after']:,.0f}，相对变化 +{d['relative_change_percent']:.2f}%。")
lines.extend(['气道：体积88.428 mL、183,975体素、26邻域66个连通块，均未改变。',
              '肺：体积8072.051 mL、16,793,887体素，逐体素未改变。','',
              '二、前后版本一致性（以补标后为参考，不是对独立真值的前后准确率）'])
for s in ['pulmonary_artery','pulmonary_vein','airway']:
    v=reports[s]['metric_sets']['vesmamba']['metrics'];a=reports[s]['metric_sets']['airrc']['metrics']
    lines.append(f"{names[s]}：Dice {100*v['binary.Dice']:.2f}%；clDice {100*v['binary.clDice']:.2f}%；NSD(1mm) {100*v['binary.NSD']:.2f}%；HD95 {v['binary.HD95_mm']:.3f} mm；Precision {100*a[s+'.Precision']:.2f}%；Recall {100*a[s+'.Recall']:.2f}%。")
lines.extend(['肺的Dice、Precision、Recall均为100%；肺未运行管状结构骨架与表面评价。',
              'HD95为0不代表所有边界不变：它是面积加权95%分位数，少量局部增删可能不改变这个分位数。','',
              '三、连通块减少的来源与新问题',
              '动脉原131个非主块中：30个的保留部分归入新主块，100个全部删除，1个仍处于非主块；另有1个纯新增孤立块（68体素）。',
              '静脉原85个非主块中：18个的保留部分归入新主块，29个全部删除，37个保留于单个非主块，1个保留部分分成两个新块；另有6个纯新增孤立块，共80体素。',
              '补标前上述结构间无重叠；补标后动脉/静脉重叠836体素（24处，0.401827 mL），气道/动脉重叠27体素（11处，0.012978 mL）。气道/静脉仍无重叠。',
              '建议结合CT核对重叠区域的类别归属，以及新孤立块和被删除片段。仅凭掩码几何不能判定每处修改是否解剖正确。','',
              '四、使用与解释',
              '查看 Case10补标前后指标对比.xlsx；results 保存逐结构JSON、汇总JSON、CSV及测试/核验记录。',
              'changes 中 *_added / *_removed 为新增/删除区域；after_*_overlap 为补标后的类别重叠区域，可作为标签掩码加载到3D Slicer并与CT叠加。',
              'inputs/before 与 inputs/after 是统一格式的独立二值通道；未解决重叠前不要按任意优先级合并成单标签。',''])
lines.extend(k+'：'+v for k,v in notes[1:])
lines.extend(['','五、复现',
              '指标环境：'+sys.executable,
              '运行本目录 rerun_metrics.ps1，可对已准备好的inputs重新调用原始指标包。',
              'prepare_comparison.py 记录原始NIfTI与NRRD的转换和结构审计；lineage_audit.py记录组件追踪。',
              'build_report.py 使用系统Python和openpyxl导出本Excel与说明。'])
(ROOT/'结果解读.txt').write_text('\n'.join(lines)+'\n',encoding='utf-8-sig')

(RES/'final_validation.json').write_text(json.dumps({'created':datetime.now().isoformat(timespec='seconds'),'all_passed':all(x['passed'] for x in checks),'checks':checks,'workbook_sha256':sha256(bookfile)},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'workbook':str(bookfile),'sheets':list(tables),'integrity_and_math_checks_passed':len(checks),'report':str(ROOT/'结果解读.txt')},ensure_ascii=False,indent=2))
