from pathlib import Path
import json,csv,statistics,hashlib,shutil,importlib.metadata
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
BASE=Path('data/airrc')
OUT=BASE/'分割效果评估_20261022'
rows=json.loads((OUT/'metrics_per_case.json').read_text(encoding='utf-8'))
assert len(rows)==12
audit=json.loads((OUT/'input_audit.json').read_text(encoding='utf-8'))
mapping=list(csv.DictReader((BASE/'原始CT与标准mask/病例对应表.csv').open(encoding='utf-8-sig')))
for a,m in zip(audit,mapping):assert a['case']==m['病例'] and a['gt_sha256']==m['mask文件SHA256']
original_root=Path('data/airrc_original_predictions')
for a in audit:
    with (original_root/a['case']/'segmentation.seg.nrrd').open('rb') as f:
        assert hashlib.file_digest(f,'sha256').hexdigest()==a['prediction_sha256']
metrics=['Dice','clDice','NSD','HD95_mm','Precision','Recall','IoU']
summary=[]
for label in range(1,5):
    group=[r for r in rows if r['label']==label]
    s={'structure':group[0]['structure'],'n':3}
    for k in metrics:
        vals=[r[k] for r in group];s[k+'_mean']=statistics.mean(vals);s[k+'_sd']=statistics.stdev(vals)
    summary.append(s)
(OUT/'metrics_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
protocol=[
'数据：AirRC三个病例，逐例比较用户目录内现有TotalSegmentator预测与原始四类label。未重新运行分割。',
'参考label保持不变，1 mm各向同性；预测按NIfTI/NRRD物理坐标最近邻重采样到label网格，网格外为0；没有平滑、碎片清理或人工修正。',
'病例CT与NRRD形状和affine一致；三例GT SHA256与既有病例对应表一致。',
'三例预测NRRD的SHA256与原始TotalSegmentator交付文件完全一致。原运行记录版本为TotalSegmentator 2.13.0。',
'GT映射：1气道腔、2气道壁、3肺动脉、4肺静脉。预测按NRRD的Segment Name/Layer/LabelValue读取独立图层。肺没有对应GT，不评分。',
'Dice=2TP/(预测体素数+GT体素数)；Precision=TP/(TP+FP)；Recall=TP/(TP+FN)；IoU=TP/(TP+FP+FN)。',
'clDice采用Lee三维骨架，按骨架精确率与敏感度的调和平均计算。气道壁是壳状结构，其clDice解释需谨慎。',
'NSD为面积加权表面Dice，容差1 mm；HD95为双向面积加权95百分位表面距离的较大值，单位mm。',
'表面元及面积沿用surface-distance 0.1；为降低内存，以精确cKDTree最近邻替代稠密距离变换，仅在表面点求距离。已用各向同性及非各向同性小体数据核对距离、HD95和多个容差NSD。',
'三例等权宏平均，样本标准差分母n-1；只有三例描述性结果，不做总体性能推断；与模型训练数据的独立性尚未核实。',
'Dice/clDice/NSD/Precision/Recall/IoU范围0–1；表格显示为百分数。HD95为mm，体积为mL。NSD容差等于本次GT网格的一个体素，薄气道壁可能Dice较低而NSD较高。',
'原始数据不修改。评估脚本和指标实现副本随结果保存。',
'标签定义来源：https://www.nature.com/articles/s41597-025-06074-6',
]
text='AirRC 三例分割效果评估\n\n各结构三例均值±样本标准差：\n'
for s in summary:
    text+=s['structure']+'：'+ '；'.join(k+' '+f"{s[k+'_mean']*(1 if k=='HD95_mm' else 100):.2f} ± {s[k+'_sd']*(1 if k=='HD95_mm' else 100):.2f}"+(' mm' if k=='HD95_mm' else '%') for k in metrics)+'\n'
text+='\n逐例Dice：\n'
for case in ['case01','case02','case03']:
    group=[r for r in rows if r['case']==case]
    text+=case+'：'+'；'.join(r['structure']+f" {r['Dice']*100:.2f}%" for r in group)+f"；四类等权平均 {statistics.mean(r['Dice'] for r in group)*100:.2f}%\n"
text+='''
结果解读：
1. 按与现有label的重叠程度，肺静脉和肺动脉表现最好，气道腔次之，气道壁是主要短板。
2. case02四类平均Dice最低（77.31%），其中气道壁55.88%，应优先在CT上复核。case01与case03整体接近：case01气道较好，case03血管较好。
3. 气道腔平均Precision 90.04%、Recall 81.24%，提示相对参考label漏分较多；仅凭汇总分数不能确定漏分位置或分支级别。
4. 气道壁平均Precision 57.57%、Recall 75.63%，同时存在标签外预测和漏分，标签外预测更突出。其NSD约96.76%不能替代Dice约65.38%的结论，1 mm容差对薄壁较宽松。
5. 动静脉平均Recall均超过91%，大部分参考前景被检出，同时仍有标签外阳性。高clDice表示骨架覆盖较好，不等于动静脉解剖分类或所有末梢均正确。
6. 本次是对提供label的一致性定量评估，没有逐层人工影像复核。不能把标签外预测直接认定为某种具体解剖误分，也不能凭三例推断总体或临床性能。
'''
text+='\n协议与解释边界：\n'+'\n'.join(protocol)
(OUT/'结果解读.txt').write_text(text,encoding='utf-8')
wb=openpyxl.Workbook();ws=wb.active;ws.title='逐例指标'
fields=[k for k in rows[0] if k!='status'];ws.append(fields)
for r in rows:ws.append([r[k] for k in fields])
ss=wb.create_sheet('三例汇总');sf=list(summary[0]);ss.append(sf)
for r in summary:ss.append([r[k] for k in sf])
notes=wb.create_sheet('评估协议')
for line in protocol:notes.append([line])
for sheet in wb:
    sheet.freeze_panes='A2'
    for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24536A')
    sheet.auto_filter.ref=sheet.dimensions
    for col in sheet.columns:
        letter=col[0].column_letter;sheet.column_dimensions[letter].width=23
        for c in list(col)[1:]:
            c.alignment=Alignment(vertical='top')
            if isinstance(c.value,float):c.number_format='0.0000'
    for j,c in enumerate(sheet[1],1):
        if isinstance(c.value,str) and any(c.value==k or c.value.startswith(k+'_') for k in metrics if k!='HD95_mm'):
            for row in sheet.iter_rows(min_row=2,min_col=j,max_col=j):row[0].number_format='0.00%'
notes.column_dimensions['A'].width=120
for row in notes:
    row[0].alignment=Alignment(wrap_text=True,vertical='top');notes.row_dimensions[row[0].row].height=45
wb.save(OUT/'AirRC三例分割效果评估.xlsx')
code=OUT/'code';code.mkdir(exist_ok=True)
for file in ['evaluate_airrc_three.py','summarize_airrc_three.py']:shutil.copy2(Path(__file__).resolve().parent/file,code/file)
shutil.copy2(Path(__file__).resolve().parent/'vesmamba_metrics.py',code/'vesmamba_metrics.py')
env={k:importlib.metadata.version(k) for k in ['numpy','scipy','scikit-image','nibabel','openpyxl']}
(OUT/'environment.json').write_text(json.dumps(env,indent=2),encoding='utf-8')
print(text)
